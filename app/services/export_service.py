"""报表导出：把账务数据导出成 Excel。

真实财务工作里，"给老板/税务局交报表"是刚需。
导出的工作簿包含三张表：账目明细、账户余额、经营概况。
"""

from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.services import stats_service

# 统一的表格样式
_HEADER_FILL = PatternFill("solid", fgColor="1E293B")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
_TITLE_FONT = Font(bold=True, size=14)
_THIN = Side(style="thin", color="D1D5DB")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_MONEY = "#,##0.00"


def _style_header(ws, row: int, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _BORDER


def _autofit(ws) -> None:
    """按内容粗略调整列宽。"""
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(width * 1.8 + 4, 40)


def build_workbook() -> bytes:
    """生成 Excel 报表，返回二进制内容。"""
    wb = Workbook()

    # ---------- 表1：账目明细 ----------
    ws = wb.active
    ws.title = "账目明细"
    ws.append(["日期", "摘要", "科目", "借方科目", "贷方科目", "金额", "类型", "凭证"])
    _style_header(ws, 1, 8)

    entries = stats_service.list_entries(limit=5000)
    type_cn = {"income": "收入", "expense": "支出", "asset": "资产", "equity": "权益",
               "transfer": "转账"}
    for e in sorted(entries, key=lambda x: x["date"]):
        ws.append([
            e["date"], e["description"], e["category"],
            e["debit_account"], e["credit_account"], e["amount"],
            type_cn.get(e["entry_type"], e["entry_type"]),
            "有" if e.get("attachment") else "",
        ])
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=6):
        for c in row:
            c.number_format = _MONEY
    _autofit(ws)

    # ---------- 表2：账户余额（资产负债表）----------
    ws2 = wb.create_sheet("账户余额")
    ws2.append(["账户余额表"])
    ws2["A1"].font = _TITLE_FONT

    bal = stats_service.account_balances()
    row = 3
    for title, items, total in [
        ("资产", bal["assets"], bal["total_assets"]),
        ("负债", bal["liabilities"], None),
        ("所有者权益", bal["equity"], bal["total_liabilities_equity"]),
    ]:
        ws2.cell(row=row, column=1, value=title).font = Font(bold=True)
        row += 1
        ws2.cell(row=row, column=1, value="科目")
        ws2.cell(row=row, column=2, value="余额（元）")
        _style_header(ws2, row, 2)
        row += 1
        for a in items:
            ws2.cell(row=row, column=1, value=a["account"])
            ws2.cell(row=row, column=2, value=a["balance"]).number_format = _MONEY
            row += 1
        if total is not None:
            ws2.cell(row=row, column=1, value=f"{title}合计").font = Font(bold=True)
            ws2.cell(row=row, column=2, value=total).number_format = _MONEY
            ws2.cell(row=row, column=2).font = Font(bold=True)
            row += 1
        row += 1

    ws2.cell(row=row, column=1, value="会计恒等式检查").font = Font(bold=True)
    row += 1
    ws2.cell(row=row, column=1, value="资产合计")
    ws2.cell(row=row, column=2, value=bal["total_assets"]).number_format = _MONEY
    row += 1
    ws2.cell(row=row, column=1, value="负债 + 所有者权益")
    ws2.cell(row=row, column=2, value=bal["total_liabilities_equity"]).number_format = _MONEY
    row += 1
    ws2.cell(row=row, column=1,
             value="结论").font = Font(bold=True)
    ws2.cell(row=row, column=2,
             value="平衡" if bal["balance_sheet_ok"] else "不平衡")
    ws2.column_dimensions["A"].width = 32
    ws2.column_dimensions["B"].width = 20

    # ---------- 表3：经营概况 ----------
    ws3 = wb.create_sheet("经营概况")
    ov = stats_service.overview()
    monthly = stats_service.monthly()

    ws3.append(["经营概况"]); ws3["A1"].font = _TITLE_FONT
    ws3.append([])
    for label, value in [
        ("账目笔数", ov["total_entries"]),
        ("收入合计", ov["total_income"]),
        ("支出合计", ov["total_expense"]),
        ("净利", ov["net"]),
        ("流动资金", bal["working_capital"]),
    ]:
        ws3.append([label, value])
    for r in range(3, ws3.max_row + 1):
        if isinstance(ws3.cell(row=r, column=2).value, (int, float)):
            ws3.cell(row=r, column=2).number_format = _MONEY

    ws3.append([])
    ws3.append(["月份", "收入", "支出", "净额"])
    head = ws3.max_row
    _style_header(ws3, head, 4)
    for i, m in enumerate(monthly["months"]):
        inc, exp = monthly["income"][i], monthly["expense"][i]
        ws3.append([m, inc, exp, round(inc - exp, 2)])
    for r in range(head + 1, ws3.max_row + 1):
        for c in range(2, 5):
            ws3.cell(row=r, column=c).number_format = _MONEY
    ws3.column_dimensions["A"].width = 16
    for col in "BCD":
        ws3.column_dimensions[col].width = 16

    # ---------- 输出 ----------
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def report_filename() -> str:
    return f"财务报表_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
